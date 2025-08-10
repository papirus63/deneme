from __future__ import annotations


def chunk_text(text: str, max_tokens: int = 800, overlap: int = 100) -> list[str]:
    """Basit karakter tabanlı bölücü."""
    size = max_tokens * 4  # ~ 1 token ≈ 4 char kabulü
    ov = overlap * 4
    out = []
    i = 0
    n = len(text)
    while i < n:
        j = min(i + size, n)
        chunk = text[i:j]
        out.append(chunk.strip())
        i = j - ov
        if i < 0:
            i = 0
    return [c for c in out if c]


def extract_text_from_upload(filename: str, data: bytes) -> str:
    """Yüklenen dosyanın içeriğini düz metin olarak döndür."""
    name = filename.lower()
    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore")
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        from io import BytesIO
        reader = PdfReader(BytesIO(data))
        return "\n".join([page.extract_text() or "" for page in reader.pages])
    if name.endswith(".docx"):
        import docx
        from io import BytesIO
        doc = docx.Document(BytesIO(data))
        return "\n".join([p.text for p in doc.paragraphs])
    if name.endswith(".xlsx"):
        import openpyxl
        from io import BytesIO
        wb = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
        texts = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                # None değerleri temizle ve birleştir
                line = " ".join(str(cell) for cell in row if cell is not None)
                if line.strip():
                    texts.append(line.strip())
        return "\n".join(texts)
    raise ValueError("Desteklenmeyen dosya türü (txt, pdf, docx, xlsx deneyin).")

